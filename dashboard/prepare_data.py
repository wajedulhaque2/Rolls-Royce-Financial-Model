"""Extract a compact, reproducible snapshot of the cached Excel valuation model."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "workbook" / "Rolls_Royce_Financial_Model_Portfolio.xlsx"
OUTPUT = Path(__file__).resolve().parent / "data" / "snapshot.json"


def extract(path: Path = WORKBOOK) -> dict:
    book = load_workbook(path, read_only=True, data_only=True)
    forecast, dcf, peers, dashboard, checks = [book[name] for name in
        ("Forecast_Model", "DCF", "Comparable_Companies", "Dashboard", "Checks")]
    years = [forecast.cell(2, col).value for col in range(2, 8)]
    series = {name: [forecast.cell(row, col).value for col in range(2, 8)] for name, row in {
        "Revenue": 9, "Operating profit": 17, "EBIT margin": 18,
        "Civil Aerospace": 5, "Defence": 6, "Power Systems": 7,
    }.items()}
    series["UFCF"] = [None, *[forecast.cell(28, col).value for col in range(3, 8)]]
    snapshot = {
        "reference_date": dcf["B4"].value.date().isoformat(),
        "reference_price": dcf["B20"].value,
        "base_wacc": dcf["B26"].value,
        "base_growth": dcf["B27"].value,
        "net_cash_m": dcf["B61"].value,
        "shares_m": dcf["B63"].value,
        "h1_ufcf_m": dcf["B38"].value,
        "terminal_period": dcf["B56"].value,
        "cash_flows_m": [dcf.cell(45, col).value for col in range(2, 7)],
        "discount_periods": [dcf.cell(46, col).value for col in range(2, 7)],
        "cached_dcf": dcf["B64"].value,
        "cached_comps": peers["B39"].value,
        "terminal_ev_share": dcf["B73"].value,
        "scenario_values": {name: dashboard.cell(10, col).value for name, col in
                            (("Bear", 9), ("Base", 10), ("Bull", 11))},
        "years": years,
        "forecast": series,
        "peers": [{"name": peers.cell(row, 1).value, "ev_revenue": peers.cell(row, 9).value,
                   "ev_ebit": peers.cell(row, 10).value} for row in range(6, 11)],
        "checks": [{"name": checks.cell(row, 1).value, "status": checks.cell(row, 3).value}
                   for row in range(6, 19)],
    }
    if len(snapshot["cash_flows_m"]) != 5 or len(snapshot["peers"]) != 5:
        raise ValueError("Workbook schema changed")
    return snapshot


if __name__ == "__main__":
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(extract(), indent=2), encoding="utf-8")
    print(f"Saved {OUTPUT}")
